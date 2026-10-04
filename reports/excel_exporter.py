import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from typing import Optional, List, Dict, Any
from repositories.policia_repository import PoliciaRepository
from repositories.mes_repository import MesRepository
from repositories.ticket_repository import TicketRepository
from repositories.vale_repository import ValeRepository
from repositories.pago_repository import PagoRepository
from repositories.venta_repository import VentaRepository
from repositories.observacion_repository import ObservacionRepository

class ExcelExporter:
    def __init__(self):
        self.policia_repo = PoliciaRepository()
        self.mes_repo = MesRepository()
        self.ticket_repo = TicketRepository()
        self.vale_repo = ValeRepository()
        self.pago_repo = PagoRepository()
        self.venta_repo = VentaRepository()
        self.obs_repo = ObservacionRepository()

    def exportar_reporte_mensual(self, anio: int, mes: int, output_path: str, local_id: str = "restaurante") -> bool:
        wb = openpyxl.Workbook()
        m = self.mes_repo.get_by_anio_mes(anio, mes)
        nombre_mes = m.nombre_mes if m else f"{mes:02d}/{anio}"

        # Setup styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="1F4E78")
        bold_font = Font(name="Arial", size=10, bold=True)
        regular_font = Font(name="Arial", size=10)
        border_thin = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # 1. RESUMEN EJECUTIVO SHEET
        ws_resumen = wb.active
        ws_resumen.title = "RESUMEN MENSUAL"
        ws_resumen.views.sheetView[0].showGridLines = True

        ws_resumen.cell(row=1, column=1, value=f"REPORTE MENSUAL - {nombre_mes} ({local_id.upper()})").font = title_font
        ws_resumen.cell(row=2, column=1, value=f"Generado automáticamente por LA PROVINCIAL").font = Font(italic=True, color="595959")

        res_ventas = self.venta_repo.get_resumen_mes(anio, mes, local_id)
        res_vales = self.vale_repo.get_resumen_mes(anio, mes, local_id)
        res_pagos = self.pago_repo.get_resumen_mes(anio, mes, local_id)
        diarios = self.ticket_repo.get_tickets_diarios_mes(anio, mes, local_id)

        tot_unidad = sum(d.para_unidad for d in diarios)
        tot_local = sum(d.local for d in diarios)
        tot_policias = sum(d.total_policias for d in diarios)

        kpis = [
            ("CONCEPTO", "VALOR", "DETALLE"),
            ("VENTA TOTAL", f"S/ {res_ventas.get('venta_total', 0.0):,.2f}", "Total ventas registradas del mes"),
            ("EFECTIVO", f"S/ {res_ventas.get('efectivo', 0.0):,.2f}", "Total efectivo cobrado"),
            ("YAPE", f"S/ {res_ventas.get('yape', 0.0):,.2f}", "Total ventas por Yape"),
            ("VENTA SIN TICKETS", f"S/ {res_ventas.get('venta_sin_tickets', 0.0):,.2f}", "Venta de clientes no policías"),
            ("VENTA TICKETS POLICIALES", f"S/ {res_ventas.get('venta_tickets', 0.0):,.2f}", f"Tickets x S/ {m.precio_ticket_policial if m else 12.0:.2f}"),
            ("PARA UNIDAD", tot_unidad, "Tickets llevados a la unidad"),
            ("LOCAL", tot_local, "Tickets consumidos en local"),
            ("TOTAL TICKETS POLICIALES", tot_policias, "Para unidad + Local"),
            ("VALES ENTREGADOS", res_vales.get('entregados', 0), "Vales emitidos"),
            ("VALES CANJEADOS", res_vales.get('canjeados', 0), "Vales canjeados"),
            ("SALDO DE VALES", res_vales.get('saldo', 0), "Entregados - Canjeados"),
            ("TICKETS PAGADOS", res_pagos.get('pagados', 0), "Tickets pagados del mes"),
            ("TICKETS DEBIDOS", res_pagos.get('debidos', 0), "Tickets pendientes de pago"),
            ("MONTO PAGADO", f"S/ {res_pagos.get('monto_pagado', 0.0):,.2f}", "Dinero cobrado por tickets"),
            ("MONTO PENDIENTE", f"S/ {res_pagos.get('monto_pendiente', 0.0):,.2f}", "Deuda por cobrar"),
        ]

        for r_idx, row in enumerate(kpis, start=4):
            for c_idx, val in enumerate(row, start=1):
                cell = ws_resumen.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border_thin
                if r_idx == 4:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.font = bold_font if c_idx == 1 else regular_font
                    if c_idx == 2:
                        cell.alignment = Alignment(horizontal="right")

        # 2. CALENDARIO POLICIAL SHEET
        ws_cal = wb.create_sheet(title="CALENDARIO POLICIAL")
        ws_cal.views.sheetView[0].showGridLines = True
        ws_cal.cell(row=1, column=1, value=f"CALENDARIO DIARIO - {nombre_mes}").font = title_font

        cal_headers = ["FECHA", "DÍA", "PARA UNIDAD", "LOCAL", "TOTAL POLICÍAS", "VALES POLICIALES", "OBSERVACIÓN"]
        for c_idx, h in enumerate(cal_headers, start=1):
            cell = ws_cal.cell(row=3, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border_thin

        for r_idx, td in enumerate(diarios, start=4):
            vals = [td.fecha, td.dia_semana, td.para_unidad, td.local, td.total_policias, td.vales_policiales, td.observacion or ""]
            for c_idx, v in enumerate(vals, start=1):
                cell = ws_cal.cell(row=r_idx, column=c_idx, value=v)
                cell.font = regular_font
                cell.border = border_thin
                if c_idx in (3, 4, 5, 6):
                    cell.alignment = Alignment(horizontal="right")

        # 3. CONTROL DE VENTAS DIARIAS SHEET
        ws_v = wb.create_sheet(title="VENTAS DIARIAS")
        ws_v.views.sheetView[0].showGridLines = True
        ws_v.cell(row=1, column=1, value=f"CONTROL DE VENTAS DIARIAS - {nombre_mes}").font = title_font

        v_headers = [
            "FECHA", "EFECTIVO", "YAPE", "VENTA SIN TICKETS", "CANT. TICKETS",
            "VENTA TICKETS", "VENTA TOTAL", "PARA UNIDAD", "LOCAL", "VALES CANJ.", "PAGADOS", "DEBIDOS", "OBSERVACIONES"
        ]
        for c_idx, h in enumerate(v_headers, start=1):
            cell = ws_v.cell(row=3, column=c_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border_thin

        ventas = self.venta_repo.get_by_mes(anio, mes, local_id)
        for r_idx, vd in enumerate(ventas, start=4):
            vals = [
                vd.fecha, vd.efectivo, vd.yape, vd.venta_sin_tickets, vd.cantidad_tickets,
                vd.venta_tickets, vd.venta_total, vd.para_unidad, vd.local,
                vd.vales_consumidos, vd.tickets_pagados, vd.tickets_debidos, vd.observaciones or ""
            ]
            for c_idx, v in enumerate(vals, start=1):
                cell = ws_v.cell(row=r_idx, column=c_idx, value=v)
                cell.font = regular_font
                cell.border = border_thin
                if c_idx >= 2 and c_idx <= 12:
                    cell.alignment = Alignment(horizontal="right")

        # 4. TICKETS POR POLICIA SHEET (Matrix)
        if m:
            ws_matrix = wb.create_sheet(title="TICKETS POR POLICIA")
            ws_matrix.views.sheetView[0].showGridLines = True
            ws_matrix.cell(row=1, column=1, value=f"MATRIZ DE TICKETS - {nombre_mes}").font = title_font

            policias = self.mes_repo.get_policias_en_mes(m.id)
            import calendar
            num_dias = calendar.monthrange(anio, mes)[1]
            matriz = self.ticket_repo.get_tickets_matriz_mes(m.id, local_id)

            m_headers = ["CÓDIGO", "APELLIDOS", "NOMBRES", "SA-PNP", "ÁREA"] + [str(d) for d in range(1, num_dias + 1)] + ["TOTAL"]
            for c_idx, h in enumerate(m_headers, start=1):
                cell = ws_matrix.cell(row=3, column=c_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
                cell.border = border_thin

            for r_idx, p in enumerate(policias, start=4):
                p_id = p["policia_id"]
                dias_map = matriz.get(p_id, {})
                dias_marks = [("X" if dias_map.get(d) == 1 else "") for d in range(1, num_dias + 1)]
                total_p = sum(1 for d in range(1, num_dias + 1) if dias_map.get(d) == 1)

                vals = [p["codigo"], p["apellidos"], p["nombres"], p["sa_pnp"] or "", p["area_historica"]] + dias_marks + [total_p]
                for c_idx, v in enumerate(vals, start=1):
                    cell = ws_matrix.cell(row=r_idx, column=c_idx, value=v)
                    cell.font = regular_font
                    cell.border = border_thin
                    if c_idx > 5:
                        cell.alignment = Alignment(horizontal="center")

        # Adjust column widths
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb.save(output_path)
        return True
